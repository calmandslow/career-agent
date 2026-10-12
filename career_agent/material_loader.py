from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable


ImageReader = Callable[[Path], str]


@dataclass(frozen=True, slots=True)
class LoadedMaterial:
    source: str
    kind: str
    text: str


class MaterialLoader:
    TEXT_SUFFIXES = {".txt", ".md"}
    IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}

    def __init__(self, image_reader: ImageReader | None = None) -> None:
        self.image_reader = image_reader

    def load(self, file_path: str | Path) -> LoadedMaterial:
        path = Path(file_path).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Material file not found: {path}")

        suffix = path.suffix.lower()
        if suffix in self.TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8-sig")
            kind = "text"
        elif suffix == ".pdf":
            text = self._read_pdf(path)
            kind = "pdf"
        elif suffix == ".docx":
            text = self._read_docx(path)
            kind = "docx"
        elif suffix == ".pptx":
            text = self._read_pptx(path)
            kind = "pptx"
        elif suffix in self.IMAGE_SUFFIXES:
            if self.image_reader is None:
                raise RuntimeError("An image-capable model is required to read screenshot materials")
            text = self.image_reader(path)
            kind = "image"
        else:
            raise ValueError(
                f"Unsupported material format: {suffix}. "
                "Supported: txt, md, pdf, docx, pptx, png, jpg, jpeg, webp"
            )

        text = text.strip()
        if not text:
            raise ValueError(f"No readable text found in material: {path.name}")
        return LoadedMaterial(source=path.name, kind=kind, text=text)

    @staticmethod
    def combine(materials: list[LoadedMaterial]) -> str:
        return "\n\n".join(
            f"[SOURCE: {material.source}]\n{material.text}" for material in materials
        )

    @staticmethod
    def _read_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install dependencies with: pip install -r requirements.txt") from exc
        reader = PdfReader(path)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    @staticmethod
    def _read_docx(path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Install dependencies with: pip install -r requirements.txt") from exc
        document = Document(path)
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(blocks)

    @staticmethod
    def _read_pptx(path: Path) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("Install dependencies with: pip install -r requirements.txt") from exc
        presentation = Presentation(path)
        blocks: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            slide_text = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            if slide_text:
                blocks.append(f"[SLIDE {index}]\n" + "\n".join(slide_text))
        return "\n\n".join(blocks)
